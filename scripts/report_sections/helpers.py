from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

NAVY_BLUE = RGBColor(30, 58, 138)       # #1E3A8A
SLATE_BLUE = RGBColor(37, 99, 235)      # #2563EB
DARK_SLATE = RGBColor(15, 23, 42)       # #0F172A
BODY_COLOR = RGBColor(51, 65, 85)       # #334155
MUTED_COLOR = RGBColor(100, 116, 139)   # #64748B

def set_cell_background(cell, hex_color):
    """Set background fill color of a table cell."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    """Set inner padding of a table cell in dxa (1/20 pt)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_callout(doc, text, title="ARCHITECTURAL PRINCIPLE / KEY TAKEAWAY"):
    """Adds a callout note panel with thick blue left border and light shading."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="36" w:space="0" w:color="2563EB"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run_title = p.add_run(f"📌 {title}\n")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(9.5)
    run_title.font.bold = True
    run_title.font.color.rgb = SLATE_BLUE
    
    run_text = p.add_run(text)
    run_text.font.name = "Calibri"
    run_text.font.size = Pt(10)
    run_text.font.italic = True
    run_text.font.color.rgb = BODY_COLOR
    doc.add_paragraph()

def add_code_block(doc, code_str, caption=None):
    """Adds a dark-themed syntax code block."""
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.paragraph_format.space_before = Pt(8)
        p_cap.paragraph_format.space_after = Pt(2)
        r_cap = p_cap.add_run(f"Listing: {caption}")
        r_cap.font.name = "Calibri"
        r_cap.font.size = Pt(9.5)
        r_cap.font.bold = True
        r_cap.font.color.rgb = RGBColor(71, 85, 105)
        
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "0F172A")
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(code_str)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(226, 232, 240)
    doc.add_paragraph()

def style_table(tbl, col_widths, headers, data):
    """Styles a formal academic table with a navy header and alternating rows."""
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    # Header Row
    hdr_cells = tbl.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1E3A8A")
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(9.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            
    # Data Rows
    for row_idx, row_data in enumerate(data):
        row_cells = tbl.add_row().cells
        bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for i, val in enumerate(row_data):
            row_cells[i].text = str(val)
            set_cell_background(row_cells[i], bg_color)
            set_cell_margins(row_cells[i], top=70, bottom=70, left=120, right=120)
            p = row_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.name = "Calibri"
                run.font.size = Pt(9.0)
                run.font.color.rgb = BODY_COLOR
                
    # Apply widths
    for row in tbl.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = Inches(width)

def add_chapter_heading(doc, text):
    """Creates a standardized Chapter heading."""
    h = doc.add_heading(text, level=1)
    h.paragraph_format.space_before = Pt(20)
    h.paragraph_format.space_after = Pt(14)
    for r in h.runs:
        r.font.name = "Calibri"
        r.font.size = Pt(18)
        r.font.bold = True
        r.font.color.rgb = NAVY_BLUE
    return h

def add_section_heading(doc, text):
    """Creates a standardized Section heading (Level 2)."""
    h = doc.add_heading(text, level=2)
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(8)
    for r in h.runs:
        r.font.name = "Calibri"
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = SLATE_BLUE
    return h

def add_sub_section_heading(doc, text):
    """Creates a standardized Sub-section heading (Level 3)."""
    h = doc.add_heading(text, level=3)
    h.paragraph_format.space_before = Pt(10)
    h.paragraph_format.space_after = Pt(4)
    for r in h.runs:
        r.font.name = "Calibri"
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = DARK_SLATE
    return h

def p(doc, text, bold_prefix=None, space_after=6):
    """Helper to add standard body paragraph with optional bold prefix."""
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(space_after)
    para.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = para.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
        r_pre.font.color.rgb = DARK_SLATE
    r_body = para.add_run(text)
    r_body.font.name = "Calibri"
    r_body.font.size = Pt(11)
    r_body.font.color.rgb = BODY_COLOR
    return para

def add_diagram_image(doc, image_path, caption, width_inches=6.2):
    """Embeds a high-resolution PNG diagram image with a centered caption."""
    import os
    if not os.path.exists(image_path):
        return
        
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(4)
    run_img = p_img.add_run()
    run_img.add_picture(image_path, width=Inches(width_inches))
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(10)
    r_cap = p_cap.add_run(caption)
    r_cap.font.name = "Calibri"
    r_cap.font.size = Pt(9.5)
    r_cap.font.bold = True
    r_cap.font.color.rgb = SLATE_BLUE

