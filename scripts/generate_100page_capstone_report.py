import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Import all section builders
from report_sections.preliminary import add_preliminary_pages
from report_sections.ch1_introduction import add_chapter_1
from report_sections.ch2_literature import add_chapter_2
from report_sections.ch3_srs_architecture import add_chapter_3
from report_sections.ch4_modules_part1 import add_chapter_4
from report_sections.ch5_modules_part2 import add_chapter_5
from report_sections.ch6_modules_part3 import add_chapter_6
from report_sections.ch7_database_design import add_chapter_7
from report_sections.ch8_algorithms import add_chapter_8
from report_sections.ch9_testing_qa import add_chapter_9
from report_sections.ch10_deployment_economics import add_chapter_10
from report_sections.ch11_results_discussion import add_chapter_11
from report_sections.ch12_conclusion_future import add_chapter_12
from report_sections.appendices_references import add_appendices_and_references

def setup_document_styling(doc):
    """Sets standard academic margins, typography, and paragraph formats."""
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        
        # Header & Footer setup
        footer = section.footer
        p_ft = footer.paragraphs[0]
        p_ft.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_ft = p_ft.add_run("SkillTrack | B.Tech Capstone Project Report — Apex Institute of Technology")
        r_ft.font.name = "Calibri"
        r_ft.font.size = Pt(8.5)
        r_ft.font.color.rgb = RGBColor(148, 163, 184)
        
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(51, 65, 85)

def main():
    print("=" * 70)
    print("SKILLTRACK 100+ PAGE CAPSTONE PROJECT REPORT GENERATOR")
    print("=" * 70)
    
    doc = Document()
    setup_document_styling(doc)
    
    sections = [
        ("Preliminary Pages (Title, Certificate, Abstract, TOC)", add_preliminary_pages),
        ("Chapter 1: Introduction & Formal Problem Statement", add_chapter_1),
        ("Chapter 2: Literature Review & Gap Analysis", add_chapter_2),
        ("Chapter 3: IEEE 830 SRS & Architecture", add_chapter_3),
        ("Chapter 4: Modules 1 to 8 (Advanced Cloud & Security Studios)", add_chapter_4),
        ("Chapter 5: Modules 9 to 16 (Developer Tooling & Duels)", add_chapter_5),
        ("Chapter 6: Modules 17 to 24 (Career, Pedagogy & AI)", add_chapter_6),
        ("Chapter 7: Database Design, Schema & Persistence Architecture", add_chapter_7),
        ("Chapter 8: Algorithmic Formulations & Mathematical Models", add_chapter_8),
        ("Chapter 9: Testing, QA & Security Hardening", add_chapter_9),
        ("Chapter 10: Serverless Deployment & Infrastructure Economics", add_chapter_10),
        ("Chapter 11: Empirical Results & Student Cohort Evaluation", add_chapter_11),
        ("Chapter 12: Conclusion & 5-Year Future Roadmap", add_chapter_12),
        ("Appendices A-D & Academic References", add_appendices_and_references)
    ]
    
    for idx, (title, builder_fn) in enumerate(sections, 1):
        print(f"[{idx:02d}/{len(sections):02d}] Generating: {title}...")
        builder_fn(doc)
        
    downloads_dir = "/Users/abhirajyadav/Downloads"
    output_download_path = os.path.join(downloads_dir, "SkillTrack_Capstone_Project_Report_100Pages.docx")
    output_workspace_path = "/Users/abhirajyadav/SkillTrack/SkillTrack_Capstone_Project_Report_100Pages.docx"
    
    print("\nSaving document...")
    doc.save(output_download_path)
    doc.save(output_workspace_path)
    
    file_size_kb = os.path.getsize(output_download_path) / 1024
    
    # Calculate document statistics
    total_paragraphs = len(doc.paragraphs)
    total_tables = len(doc.tables)
    total_words = sum(len(p.text.split()) for p in doc.paragraphs)
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                total_words += sum(len(p.text.split()) for p in c.paragraphs)
                
    print("=" * 70)
    print("DOCUMENT GENERATION COMPLETE!")
    print(f"• Download Path:    {output_download_path}")
    print(f"• Workspace Backup: {output_workspace_path}")
    print(f"• File Size:        {file_size_kb:.1f} KB")
    print(f"• Total Paragraphs: {total_paragraphs}")
    print(f"• Total Tables:     {total_tables}")
    print(f"• Total Word Count: {total_words:,} words")
    print("=" * 70)

if __name__ == "__main__":
    main()
